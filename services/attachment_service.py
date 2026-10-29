"""Resolve workspace and uploaded sources into bounded text or native visual inputs.

Label supplied content for model input while keeping binary payloads out of stored transcripts.
"""

import base64
import hashlib
import json
import warnings
from collections.abc import Generator
from contextlib import closing
from dataclasses import dataclass, field
from io import BytesIO
from math import ceil
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._reader import WorkSheetParser
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from models.app_settings import DEFAULT_APP_SETTINGS
from models.attachment import AttachmentSource, WorkspaceAttachment
from services.attachment_errors import AttachmentError
from services.attachment_store import AttachmentStore, StagedUpload
from services.bounded_io import read_bounded
from services.local_work import check_local_cancelled
from services.workspace_service import WorkspaceAccessError, WorkspaceService


@dataclass
class PreparedAttachments:
    """Carry transient provider parts and local estimates independently of metadata."""

    content: list[dict] = field(default_factory=list)
    metadata: list[dict] = field(default_factory=list)
    visual_tokens: int = 0


class AttachmentService:
    """Read supported local files once, enforcing root, size, and parser limits."""

    # Defaults below are provisional desktop resource policies.
    # Workload samples do not establish safe bounds across hardware.
    # Model token capacity does not determine these parser-work limits.
    # Count every local file independently of provider image capacity.
    MAX_FILES = DEFAULT_APP_SETTINGS.attachment_max_files
    # Keep a provisional mixed-format raw-read budget of about 50 MB.
    # The common total is a provisional budget for reading mixed formats.
    # Base64 and JSON duplicate payloads; this is not a peak-memory bound.
    MAX_FILE_BYTES = DEFAULT_APP_SETTINGS.attachment_max_file_bytes
    MAX_TOTAL_BYTES = DEFAULT_APP_SETTINGS.attachment_max_total_bytes
    # Bound text retained for tokenization and prompts, not model input tokens.
    # Parser allocations and raw input buffers have separate costs.
    MAX_TEXT_BYTES = DEFAULT_APP_SETTINGS.attachment_max_text_bytes
    # Bound Office ZIP expansion independently of compressed source size.
    # Byte and entry budgets are provisional parser-work limits.
    MAX_ZIP_BYTES = DEFAULT_APP_SETTINGS.attachment_max_zip_bytes
    MAX_ZIP_MEMBERS = DEFAULT_APP_SETTINGS.attachment_max_zip_members
    # 50 MP RGB/RGBA pixels alone need about 150-200 MB before decoder copies.
    # This provisional pixel guard is not a measured peak-memory guarantee.
    MAX_IMAGE_PIXELS = DEFAULT_APP_SETTINGS.attachment_max_image_pixels
    # A generous initial page count bounds local page-object traversal.
    # This is not an Azure page limit or benchmarked latency threshold.
    MAX_PDF_PAGES = DEFAULT_APP_SETTINGS.attachment_max_pdf_pages
    # Retain the existing initial work budget while counting stored cells.
    # Empty rectangular gaps do not count toward the provisional cell limit.
    MAX_SHEET_CELLS = DEFAULT_APP_SETTINGS.attachment_max_sheet_cells
    IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
    IMAGE_MIMES = {
        "PNG": "image/png",
        "JPEG": "image/jpeg",
        "WEBP": "image/webp",
        "GIF": "image/gif",
    }
    OFFICE_EXTENSIONS = {".docx", ".xlsx", ".pptx"}
    UNSUPPORTED_EXTENSIONS = {
        ".doc",
        ".xls",
        ".ppt",
        ".docm",
        ".xlsm",
        ".pptm",
        ".zip",
        ".7z",
        ".rar",
        ".exe",
        ".dll",
        ".bin",
        ".mp3",
        ".wav",
        ".mp4",
        ".mov",
        ".avi",
        ".heic",
        ".tiff",
    }

    def __init__(
        self, workspace: WorkspaceService, *, limits: dict | None = None,
        store: AttachmentStore | None = None,
    ) -> None:
        """Use workspace boundaries and configurable local parser resource limits."""
        self.workspace = workspace
        self.store = store
        for name, value in (limits or {}).items():
            if not name.startswith("MAX_") or not hasattr(self, name):
                raise ValueError("Unknown attachment resource limit.")
            if type(value) is not int or value < 1:
                raise ValueError("Attachment resource limits must be positive integers.")
            setattr(self, name, value)
        if Image.MAX_IMAGE_PIXELS is not None:
            self.MAX_IMAGE_PIXELS = min(
                self.MAX_IMAGE_PIXELS, Image.MAX_IMAGE_PIXELS
            )

    def prepare(self, paths: list[str]) -> PreparedAttachments:
        """Preserve the legacy workspace-relative selection interface."""
        return self.prepare_sources([
            WorkspaceAttachment(source="workspace", path=path) for path in paths
        ])

    def prepare_sources(self, sources: list[AttachmentSource]) -> PreparedAttachments:
        """Resolve source permissions separately while sharing all parser and read limits."""
        upload_ids = [
            str(source.attachment_id) for source in sources if source.source == "upload"
        ]
        if upload_ids and self.store is None:
            raise AttachmentError(
                "attachment_source_missing", "Uploaded attachment storage is unavailable. Select the file again.", "", 410,
            )
        release = self.store.claim(upload_ids) if upload_ids else None
        try:
            return self._prepare_sources(sources)
        finally:
            if release is not None:
                release()

    def _prepare_sources(self, sources: list[AttachmentSource]) -> PreparedAttachments:
        """Prepare an all-or-nothing selection against one captured workspace identity."""
        result = PreparedAttachments()
        root = self.workspace.root
        seen: set[tuple[str, object]] = set()
        total = 0
        workspace_used = False
        for source in sources:
            check_local_cancelled()
            supplied_path = source.path if source.source == "workspace" else str(source.attachment_id)
            try:
                expected = None
                if source.source == "workspace":
                    workspace_used = True
                    if (
                        source.workspace_revision is not None
                        and source.workspace_revision != self.workspace.revision
                        or source.workspace_session_id is not None
                        and source.workspace_session_id != self.workspace.session_id
                    ):
                        raise AttachmentError(
                            "attachment_workspace_changed", "Workspace selection changed. Select the file again.", supplied_path, 409,
                        )
                    path = self.workspace._resolve(source.path, root=root)
                    name = path.relative_to(root).as_posix()
                    key = ("workspace", path)
                    reference = {
                        "source": "workspace", "path": name,
                        "workspace_revision": self.workspace.revision,
                        "workspace_session_id": self.workspace.session_id,
                    }
                else:
                    expected = self.store.get(str(source.attachment_id))
                    path = self.store.root / f"{source.attachment_id}.blob"
                    name = expected["filename"]
                    key = ("upload", source.attachment_id)
                    reference = expected["reference"]
                if key in seen:
                    continue
                if len(seen) >= self.MAX_FILES:
                    raise AttachmentError(
                        "attachment_limit", f"Selection exceeds the local limit of {self.MAX_FILES} distinct attachments.", supplied_path, 413,
                    )
                data = self._read_file(path, name, self.MAX_TOTAL_BYTES - total)
                if expected is not None and (
                    len(data) != expected["size_bytes"]
                    or hashlib.sha256(data).hexdigest() != expected["sha256"]
                ):
                    raise AttachmentError(
                        "attachment_source_damaged", "The uploaded original changed or is damaged. Select the file again.", supplied_path, 410,
                    )
                total += len(data)
                parts, metadata, estimate = self._prepare_file(name, data)
                metadata.update(source=source.source, reference=reference)
                parts = self._file_content(
                    len(result.metadata) + 1, name, source.source,
                    metadata["delivery"], parts,
                )
                check_local_cancelled()
            except AttachmentError:
                raise
            except PermissionError as exc:
                raise AttachmentError(
                    "attachment_permission_denied", "Cannot read this attachment.", supplied_path, 403,
                ) from exc
            except WorkspaceAccessError as exc:
                raise AttachmentError(
                    "attachment_path_invalid", str(exc), supplied_path,
                ) from exc
            except OSError as exc:
                raise AttachmentError(
                    "attachment_unavailable", "Attachment could not be read. Select it again.", supplied_path,
                ) from exc
            seen.add(key)
            result.content.extend(parts)
            result.metadata.append(metadata)
            result.visual_tokens += estimate
        if workspace_used and self.workspace.root != root:
            raise AttachmentError(
                "attachment_workspace_changed", "Workspace changed while reading attachments. Select the files again.", "", 409,
            )
        return result

    def _read_file(self, path: Path, name: str, remaining: int) -> bytes:
        """Bound a source read even if the original grows after its size was checked."""
        if not path.is_file():
            raise AttachmentError("attachment_not_found", "Attachment file not found.", name, 404)
        if path.stat().st_size > self.MAX_FILE_BYTES:
            raise AttachmentError(
                "attachment_too_large", f"Attachment exceeds the local read limit of {self.MAX_FILE_BYTES} bytes.", name, 413,
            )
        with path.open("rb") as original:
            data = read_bounded(original, min(self.MAX_FILE_BYTES, remaining))
        if len(data) > self.MAX_FILE_BYTES or len(data) > remaining:
            raise AttachmentError(
                "attachment_too_large", f"Attachments exceed the local combined read limit of {self.MAX_TOTAL_BYTES} bytes.", name, 413,
            )
        return data

    def validate_upload(self, upload: StagedUpload) -> dict:
        """Validate staged bytes with the same format and parser policy used by chat."""
        data = self._read_file(upload.path, upload.filename, self.MAX_TOTAL_BYTES)
        _, metadata, _ = self._prepare_file(upload.filename, data)
        check_local_cancelled()
        return metadata

    def _prepare_file(
        self, path: str, data: bytes
    ) -> tuple[list[dict], dict, int]:
        """Dispatch validated bytes without trusting the filename as proof of format."""
        check_local_cancelled()
        extension = Path(path).suffix.lower()
        metadata = {"path": path, "size_bytes": len(data), "warnings": []}
        try:
            if extension in self.IMAGE_EXTENSIONS:
                part, dimensions = self._image(path, data)
                metadata.update(kind="image", delivery="native", **dimensions)
                return (
                    [part],
                    metadata,
                    self._visual_estimate(
                        dimensions["width"], dimensions["height"]
                    ),
                )
            if extension == ".pdf":
                part, pages, estimate = self._pdf(path, data)
                metadata.update(kind="pdf", delivery="native", pages=pages)
                metadata["warnings"] = [
                    "The local PDF estimate covers page visuals only; provider-extracted text and rendering can change token usage."
                ]
                return (
                    [part],
                    metadata,
                    estimate,
                )
            if extension in self.UNSUPPORTED_EXTENSIONS:
                raise AttachmentError(
                    "attachment_type_unsupported",
                    "This file format is not supported.",
                    path,
                    415,
                )
            if extension in self.OFFICE_EXTENSIONS:
                self._validate_zip(path, data)
                if extension == ".docx":
                    sections = self._word(data)
                elif extension == ".xlsx":
                    sections = self._spreadsheet(path, data)
                else:
                    sections = self._presentation(data)
                text = self._bounded_text(path, sections)
                metadata.update(kind=extension[1:], delivery="extracted_text")
                metadata["warnings"] = [
                    "Only document text, tables, and cell formulas are "
                    "included; embedded images, charts, and formatting "
                    "are omitted."
                ]
                if not text.strip():
                    raise AttachmentError(
                        "attachment_empty",
                        "No readable document text was found. Export visual "
                        "content as PDF or an image.",
                        path,
                    )
            else:
                text = data.decode("utf-8-sig")
                if any(
                    ord(character) < 32 and character not in "\t\n\r"
                    for character in text
                ):
                    raise AttachmentError(
                        "attachment_type_unsupported",
                        "File is binary rather than UTF-8 text.",
                        path,
                        415,
                    )
                metadata.update(kind="text", delivery="text")
                self._check_text(path, text)
            return (
                [
                    {
                        "type": "input_text",
                        "text": text,
                    }
                ],
                metadata,
                0,
            )
        except AttachmentError:
            raise
        except UnicodeDecodeError as exc:
            raise AttachmentError(
                "attachment_type_unsupported",
                "Use UTF-8 text, PDF, DOCX, XLSX, PPTX, or a supported image.",
                path,
                415,
            ) from exc
        except (
            Image.DecompressionBombWarning, Image.DecompressionBombError
        ) as exc:
            raise AttachmentError(
                "attachment_too_large",
                "Image exceeds the local decoder pixel limit.",
                path,
                413,
            ) from exc
        except Exception as exc:
            raise AttachmentError(
                "attachment_invalid",
                "File is damaged, encrypted, or does not match its format.",
                path,
            ) from exc

    @staticmethod
    def _file_content(
        number: int, filename: str, source: str, delivery: str, parts: list[dict],
    ) -> list[dict]:
        """Bind file identity to escaped text or its following native content part."""
        description = {
            "attachment": number, "filename": filename,
            "source": source, "delivery": delivery,
        }
        if delivery in {"text", "extracted_text"}:
            description["content"] = parts[0]["text"]
            heading = "\n\nAttached file content (JSON):\n"
            content = []
        else:
            heading = "\n\nAttached file (next content part):\n"
            content = parts
        return [{
            "type": "input_text",
            "text": heading + json.dumps(description, ensure_ascii=False),
        }, *content]

    def _check_text(self, path: str, text: str) -> None:
        """Reject oversized extracted content without silently truncating documents."""
        if len(text.encode("utf-8")) > self.MAX_TEXT_BYTES:
            raise AttachmentError(
                "attachment_text_too_large",
                f"Extracted text exceeds the local limit of {self.MAX_TEXT_BYTES} bytes.",
                path,
                413,
            )

    def _bounded_text(
        self, path: str, sections: Generator[str, None, None]
    ) -> str:
        """Stop extraction before retaining over-budget text and close parser resources."""
        parts: list[str] = []
        size = 0
        with closing(sections):
            for text in sections:
                check_local_cancelled()
                size += len(text.encode("utf-8")) + bool(parts)
                if size > self.MAX_TEXT_BYTES:
                    raise AttachmentError(
                        "attachment_text_too_large",
                        f"Extracted text exceeds the local limit of {self.MAX_TEXT_BYTES} bytes.",
                        path,
                        413,
                    )
                parts.append(text)
        return "\n".join(parts)

    def _image(self, path: str, data: bytes) -> tuple[dict, dict]:
        """Verify supported image data and reject animation or excessive dimensions."""
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(BytesIO(data)) as image:
                mime = self.IMAGE_MIMES.get(image.format)
                if not mime:
                    raise AttachmentError(
                        "attachment_type_unsupported",
                        "Unsupported image encoding.",
                        path,
                        415,
                    )
                width, height = image.size
                if width * height > self.MAX_IMAGE_PIXELS:
                    raise AttachmentError(
                        "attachment_too_large",
                        f"Image exceeds the local decode limit of {self.MAX_IMAGE_PIXELS} pixels.",
                        path,
                        413,
                    )
                if getattr(image, "is_animated", False):
                    raise AttachmentError(
                        "attachment_type_unsupported",
                        "Animated images are not supported. Select a still image.",
                        path,
                        415,
                    )
                image.verify()
            check_local_cancelled()
            with Image.open(BytesIO(data)) as image:
                image.load()
            check_local_cancelled()
        return (
            {
                "type": "input_image",
                "image_url": self._data_url(mime, data),
            },
            {"mime_type": mime, "width": width, "height": height},
        )

    def _pdf(self, path: str, data: bytes) -> tuple[dict, int, int]:
        """Inspect PDF structure without making local text extraction a native-input requirement."""
        if b"%PDF-" not in data[:1024]:
            raise AttachmentError(
                "attachment_invalid", "File is not a valid PDF.", path
            )
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise AttachmentError(
                "attachment_encrypted",
                "Password-protected PDFs are not supported. Provide an unlocked copy.",
                path,
            )
        pages = len(reader.pages)
        if not 1 <= pages <= self.MAX_PDF_PAGES:
            raise AttachmentError(
                "attachment_too_large",
                f"PDF exceeds the local page inspection limit of {self.MAX_PDF_PAGES}, or has no pages.",
                path,
                413,
            )
        estimate = 0
        for page in reader.pages:
            check_local_cancelled()
            estimate += self._visual_estimate(
                float(page.mediabox.width) * 2,
                float(page.mediabox.height) * 2,
            )
        return (
            {
                "type": "input_file",
                "filename": Path(path).name,
                "file_data": self._data_url("application/pdf", data),
            },
            pages,
            estimate,
        )

    @staticmethod
    def _visual_estimate(width: float, height: float) -> int:
        """Use dimension-based patches as an uncertain proxy, never exact model tokenization."""
        return max(1, ceil(abs(width) / 32) * ceil(abs(height) / 32))

    @staticmethod
    def _data_url(mime: str, data: bytes) -> str:
        """Encode validated bytes inline without uploading persistent provider files."""
        return f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"

    def _validate_zip(self, path: str, data: bytes) -> None:
        """Reject encrypted or oversized Office containers before invoking their parsers."""
        if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            raise AttachmentError(
                "attachment_encrypted",
                "Encrypted or legacy Office containers are not supported. "
                "Export an unlocked DOCX, XLSX, or PPTX copy.",
                path,
            )
        try:
            with ZipFile(BytesIO(data)) as archive:
                members = archive.infolist()
                if (
                    len(members) > self.MAX_ZIP_MEMBERS
                    or sum(member.file_size for member in members)
                    > self.MAX_ZIP_BYTES
                ):
                    raise AttachmentError(
                        "attachment_too_large",
                        "Expanded document exceeds the parser limit.",
                        path,
                        413,
                    )
                if any(member.flag_bits & 1 for member in members):
                    raise AttachmentError(
                        "attachment_encrypted",
                        "Encrypted documents are not supported.",
                        path,
                    )
        except BadZipFile as exc:
            raise AttachmentError(
                "attachment_invalid",
                "File is not a valid Office document.",
                path,
            ) from exc

    @staticmethod
    def _word(data: bytes) -> Generator[str, None, None]:
        """Extract ordered Word body paragraphs, tables, headers, and footers."""
        document = Document(BytesIO(data))
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                yield from (
                    "\t".join(cell.text for cell in row.cells)
                    for row in block.rows
                )
            else:
                yield block.text
        seen = set()
        for section in document.sections:
            check_local_cancelled()
            for region in (section.header, section.footer):
                if region.part.partname in seen:
                    continue
                seen.add(region.part.partname)
                yield from (
                    paragraph.text for paragraph in region.paragraphs
                )
                for table in region.tables:
                    yield from (
                        "\t".join(cell.text for cell in row.cells)
                        for row in table.rows
                    )

    def _spreadsheet(
        self, path: str, data: bytes
    ) -> Generator[str, None, None]:
        """Extract labelled spreadsheet cells and formulas without running workbook code."""
        workbook = load_workbook(
            BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        cells = 0
        try:
            for sheet in workbook.worksheets:
                check_local_cancelled()
                yield f"Sheet: {sheet.title}"
                # The public row iterator pads gaps; inspect stored XML cells.
                with sheet._get_source() as source:
                    parser = WorkSheetParser(
                        source, sheet._shared_strings, data_only=False,
                        epoch=workbook.epoch,
                        date_formats=workbook._date_formats,
                        timedelta_formats=workbook._timedelta_formats,
                    )
                    for _, row in parser.parse():
                        check_local_cancelled()
                        cells += len(row)
                        if cells > self.MAX_SHEET_CELLS:
                            raise AttachmentError(
                                "attachment_too_large",
                                f"Spreadsheet exceeds the local limit of {self.MAX_SHEET_CELLS} stored cells.",
                                path, 413,
                            )
                        values = [
                            f"{get_column_letter(cell['column'])}{cell['row']}={cell['value']}"
                            for cell in row if cell["value"] is not None
                        ]
                        if values:
                            yield "\t".join(values)
        finally:
            workbook.close()

    @staticmethod
    def _presentation(data: bytes) -> Generator[str, None, None]:
        """Extract slide text, tables, grouped shapes, and existing speaker notes."""
        presentation = Presentation(BytesIO(data))

        def shape_text(shapes) -> Generator[str, None, None]:
            """Visit nested groups while preserving the slide's shape order."""
            for shape in shapes:
                check_local_cancelled()
                if hasattr(shape, "shapes"):
                    yield from shape_text(shape.shapes)
                if shape.has_text_frame:
                    yield shape.text_frame.text
                if shape.has_table:
                    yield from (
                        "\t".join(cell.text for cell in row.cells)
                        for row in shape.table.rows
                    )

        for number, slide in enumerate(presentation.slides, 1):
            check_local_cancelled()
            yield f"Slide: {number}"
            yield from shape_text(slide.shapes)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                yield "Notes: " + slide.notes_slide.notes_text_frame.text
